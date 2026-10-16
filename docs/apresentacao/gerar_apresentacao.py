"""Generate an editable physician-facing deck. Run with PYTHONPATH=.tools/presentation."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'apresentacao'
OUT.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG, WHITE, INK, MUTED, GREEN, SOFT, LINE, GOLD = 'EFF4F2', 'FFFFFF', '0E1F1B', '52655F', '0B7A66', 'D3EBE4', 'D5E0DC', 'B16A0B'

def box(s,x,y,w,h,color=WHITE,radius=True,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: sh.adjustments[0]=.12
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color)
    sh.line.fill.background() if line is None else None
    if line: sh.line.color.rgb=RGBColor.from_string(line)
    return sh

def text(s,content,x,y,w,h,size=20,color=INK,bold=False,font='Manrope'):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01); tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(9)
    return sh

def base(kicker,title,sub=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else BG)
    text(s,kicker.upper(),.55,.34,11,.3,10,SOFT if dark else GREEN,True)
    text(s,title,.55,.85,12.1,1.0,32,WHITE if dark else INK,True,'Bricolage Grotesque')
    if sub:text(s,sub,.55,1.85,12.1,.7,15,SOFT if dark else MUTED)
    text(s,'REPASSE MED  /  Demonstração com dados fictícios · Out/2026',.55,7.12,10,.2,9,SOFT if dark else MUTED)
    text(s,f'{len(prs.slides):02}',12.2,7.1,.55,.25,10,SOFT if dark else MUTED,True)
    return s

def card(s,x,y,w,title,body,tag=None):
    box(s,x,y,w,3.0,line=LINE)
    if tag:text(s,tag,x+.22,y+.2,w-.44,.35,11,GREEN,True)
    text(s,title,x+.22,y+.7,w-.44,.7,23,INK,True,'Bricolage Grotesque')
    text(s,body,x+.22,y+1.5,w-.44,1.3,16,MUTED)

def notes(s,body):s.notes_slide.notes_text_frame.text=body

def screenshot(s,path,x,y,w,h):
    iw,ih=Image.open(path).size
    p=s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    source_ratio=iw/ih; frame_ratio=w/h
    if source_ratio<frame_ratio:p.crop_bottom=1-source_ratio/frame_ratio
    else:
        crop=(1-frame_ratio/source_ratio)/2;p.crop_left=crop;p.crop_right=crop
    return p

# 01 — cover
s=base('Clareza para quem cuida','Seu trabalho em números.\nSeu repasse com mais clareza.',dark=True)
text(s,'Repasse MED',.6,3.15,7,1,48,WHITE,True,'Bricolage Grotesque')
text(s,'Acompanhe a produção ambulatorial, entenda a previsão\ne identifique o que merece atenção.',.6,4.35,8,1.1,21,SOFT)
box(s,.6,6,6.2,.52,GREEN)
text(s,'APRESENTAÇÃO DO MVP  •  VALIDAÇÃO COM MÉDICOS',.8,6.14,5.8,.22,11,WHITE,True)
box(s,10.15,2.1,2.15,3.9,GREEN)
text(s,'PREVER\nCONFERIR\nPLANEJAR',10.4,2.7,1.7,2.8,23,WHITE,True,'Bricolage Grotesque')
notes(s,'Apresentar como demonstração funcional com dados fictícios. Não há conexão com o MV nem cálculo oficial de pagamento. Duração sugerida: 8–12 minutos nos slides 1–12. Slides seguintes são apêndices para discussão comercial e de implantação.')

# 02
s=base('O valor para o médico','Mais clareza em três momentos do mês.','Uma visão pessoal da produção, dos valores previstos e dos pontos de atenção.')
card(s,.55,2.8,3.95,'Antes de receber','Veja a previsão da competência e acompanhe sua meta.','01 / PLANEJAMENTO')
card(s,4.68,2.8,3.95,'Ao conferir','Entenda a composição por convênio, regra e atendimento.','02 / TRANSPARÊNCIA')
card(s,8.81,2.8,3.95,'Ao agir','Identifique glosas e consultas sem regra para discutir com o hospital.','03 / FOCO')
text(s,'Benefícios esperados: menos conferência dispersa e mais previsibilidade. Serão medidos no piloto.',.65,6.25,11.8,.5,14,MUTED)
notes(s,'Não prometer aumento de renda, recuperação automática de glosas ou economia de tempo já comprovada. O valor do produto é tornar a composição visível e orientar a conferência.')

# 03
s=base('Como funciona','Da entrada à conferência, em seis passos.','No dia a dia, o médico acompanha os números; o hospital confirma os dados e as regras.')
items=[('1','Entre na conta','E-mail e senha. Na demo local, o botão Entrar abre dados fictícios.'),('2','Complete seu perfil','CRM e código de prestador. Vínculo com o hospital ainda será validado.'),('3','Receba os dados','Hoje: seed ou importação CSV/XLSX. Integração MV: etapa futura.'),('4','Confira as regras','Valor fixo ou percentual, vigência e dedução por convênio.'),('5','Escolha a competência','Veja a previsão, o que já foi pago e a evolução da meta.'),('6','Investigue os detalhes','Filtre atendimentos e confira glosas ou valores sem regra.')]
for i,(n,t,b) in enumerate(items):
    x=.55+(i%3)*4.15;y=2.65+(i//3)*1.95
    box(s,x,y,3.95,1.72,WHITE,line=LINE);text(s,n,x+.18,y+.17,.4,.5,24,GREEN,True);text(s,t,x+.72,y+.17,3,.45,19,INK,True);text(s,b,x+.2,y+.79,3.55,.76,13,MUTED)
notes(s,'No piloto, a configuração contratual deve ser conferida com o financeiro do hospital. Não apresentar o CRM digitado como vínculo de identidade já verificado. Login real Supabase implementado, mas a validação do stack completo está pendente.')

# 04
s=base('Painel principal','O essencial aparece primeiro.','Captura real do MVP · novembro de 2026 · dados fictícios.')
screenshot(s,ROOT/'web/test-results/dashboard-desktop.png',.55,2.65,8.3,4.08)
text(s,'01  A receber',9.15,2.8,3.4,.4,19,GREEN,True)
text(s,'Previsão líquida ainda não paga na competência.',9.15,3.3,3.4,.7,15,MUTED)
text(s,'02  Meta e indicadores',9.15,4.25,3.4,.4,19,GREEN,True)
text(s,'Produção, ticket médio e valores recebidos.',9.15,4.75,3.4,.7,15,MUTED)
text(s,'03  Alertas acionáveis',9.15,5.65,3.4,.4,19,GREEN,True)
text(s,'Consultas sem regra e glosas para conferir.',9.15,6.1,3.4,.65,15,MUTED)
notes(s,'A captura mostra R$ 7.413,00 de previsão em novembro. São números fictícios, sem significado contratual. Se usar a demo ao vivo, mudanças de regras/importação podem alterar os valores.')

# 05
s=base('Leia os números corretamente','“A receber” é uma previsão.','O valor ajuda a planejar; a conferência oficial continua com o hospital.')
card(s,.55,2.75,3.95,'A receber','Soma líquida dos atendimentos ainda não pagos. Inclui agendados.','PREVISÃO')
card(s,4.68,2.75,3.95,'Já recebido','Soma dos atendimentos com status pago. Aparece separada da previsão.','REALIZADO FINANCEIRO')
card(s,8.81,2.75,3.95,'Competência','Na importação atual, corresponde ao mês do atendimento. O hospital validará o critério definitivo.','MÊS DE REFERÊNCIA')
box(s,.55,6.05,12.2,.55,SOFT);text(s,'Atendimento agendado pode não acontecer. Uma consulta sem regra fica fora da soma e gera alerta.',.75,6.2,11.8,.25,13,GREEN,True)
notes(s,'A barra de meta usa previsto + pago. Glosa parcial mantém o saldo líquido na previsão; essa semântica depende da validação do hospital. Não prometer data de depósito nem valor garantido.')

# 06
s=base('Por convênio e atendimento','Entenda de onde vem cada valor.','A visão consolidada dá contexto. O detalhe permite conferir a composição.')
card(s,.55,2.8,5.98,'Por convênio','Quantidade de consultas, valor de referência, honorário, subtotal e ticket médio.','VISÃO CONSOLIDADA')
card(s,6.75,2.8,5.98,'Por atendimento','Código, data, convênio, status, origem do dado e regra utilizada. Filtros por convênio, status e dia.','CONFERÊNCIA PONTUAL')
text(s,'Se houver dedução, glosa ou mudança de regra, o subtotal considera cada atendimento individualmente.',.65,6.18,11.8,.5,15,MUTED)
notes(s,'Mostre a navegação Por convênio e depois Atendimentos. Abra o detalhe de uma consulta. Não há nome/CPF do paciente: o identificador de conferência é o código do atendimento.')

# 07
s=base('Exemplo didático','Como o repasse é calculado.','Dois modelos de regra: valor fixo em reais ou percentual sobre o valor da consulta.')
box(s,.55,2.75,7.7,3.72,GREEN)
text(s,'CONSULTA DE R$ 200,00  ×  HONORÁRIO DE 60%',.85,3.05,7,.5,17,WHITE,True)
text(s,'R$ 120,00',.85,3.75,6,1,42,WHITE,True,'Bricolage Grotesque')
text(s,'− dedução de 10%  →  R$ 108,00\n− glosa parcial de 50%  →  R$ 54,00 líquidos',.85,5.05,7,1.1,21,WHITE)
text(s,'Regra fixa',8.65,2.95,3.7,.5,23,GREEN,True)
text(s,'Substitui o percentual por um honorário em reais.',8.65,3.55,3.7,.9,17,MUTED)
text(s,'Vigência',8.65,4.75,3.7,.5,23,GREEN,True)
text(s,'A regra deve estar válida na data do atendimento.',8.65,5.35,3.7,.9,17,MUTED)
notes(s,'Exemplo hipotético, sem relação com contrato real. Dedução é aplicada antes da glosa; arredondamento por atendimento em centavos. Glosa integral resulta em zero. Sem regra resulta em valor pendente, não em repasse confirmado de zero. O cálculo está centralizado em SQL.')

# 08
s=base('Ciclo de recebimento','Status ajudam a orientar a conferência.','O gráfico mostra quantidades por etapa; a legenda também informa valores líquidos.')
status=[('Agendado','Consulta prevista. Ainda pode mudar.','9AABA6'),('Realizado','Consulta registrada como realizada.','7863A3'),('A faturar','Pendente de faturamento.','C98512'),('Faturado','Faturamento registrado.','2F6FB5'),('Glosado','Glosa total ou parcial para conferir.','C6423A'),('Pago','Repasse marcado como recebido.','0B7A66')]
for i,(t,b,c) in enumerate(status):
    x=.55+(i%3)*4.15;y=2.7+(i//3)*1.85
    box(s,x,y,3.95,1.6,WHITE,line=LINE);box(s,x+.17,y+.2,.1,.4,c,False);text(s,t,x+.43,y+.2,3.25,.45,21,INK,True);text(s,b,x+.2,y+.85,3.55,.6,14,MUTED)
notes(s,'Os status atuais vêm dos dados importados ou fictícios. A fonte de cada status e sua atualização automática serão mapeadas com o hospital na integração. Não sugerir que o MVP já confirma liquidação bancária.')

# 09
s=base('Entrada de dados e cuidado com a informação','Importe apenas o necessário.','O MVP já aceita CSV/XLSX; a sincronização automática com o MV ainda não foi implementada.')
card(s,.55,2.8,5.98,'Fluxo de importação','Escolher arquivo → mapear colunas → validar → importar → conferir a previsão recalculada.','DISPONÍVEL NO MVP')
card(s,6.75,2.8,5.98,'Campos financeiros','Código de atendimento, convênio, data, status e valor opcional da consulta. Sem nome ou CPF de paciente.','DADOS MÍNIMOS')
text(s,'Até 1.000 linhas por arquivo. O acesso deve ser individual; permissões e operação real exigem validação.',.65,6.2,11.9,.6,14,MUTED)
notes(s,'Arquivos originais são lidos no navegador; somente campos permitidos são enviados. Cabeçalhos com indicação de dados pessoais são rejeitados. Isso não equivale a certificação LGPD. O código de atendimento ainda requer proteção e governança. RLS foi testado no PostgreSQL; Auth nativo e operação hospedada continuam pendentes.')

# 10
s=base('No celular ou no computador','Uma experiência para a rotina do médico.','Interface responsiva, com navegação adaptada ao tamanho da tela.')
box(s,.8,2.55,3.0,4.35,INK)
screenshot(s,ROOT/'web/test-results/dashboard-mobile.png',.93,2.68,2.74,4.07)
text(s,'Celular',4.4,2.85,7.5,.6,30,GREEN,True,'Bricolage Grotesque')
text(s,'Resumo financeiro e navegação inferior\npara consultar em poucos toques.',4.4,3.65,7.5,1.0,22,INK)
text(s,'Computador',4.4,5.05,7.5,.6,30,GREEN,True,'Bricolage Grotesque')
text(s,'Mais espaço para comparar convênios\ne investigar os atendimentos.',4.4,5.85,7.5,1.0,22,INK)
notes(s,'PWA com manifest e service worker implementados. Instalação depende de HTTPS/localhost e navegador compatível e ainda não foi validada em dispositivo físico. Sem conexão, a aplicação exige reconectar para consultar dados; não prometer consulta financeira offline.')

# 11
s=base('Proposta de planos · sujeita à validação','Começar simples. Evoluir com a necessidade.','Preços sugeridos para lançamento. No MVP, todos os recursos estão liberados e não há cobrança.')
card(s,.55,2.75,5.98,'Free  ·  R$ 0','Proposta: 1 convênio, competência atual e visão básica da produção.','PARA CONHECER')
card(s,6.75,2.75,5.98,'Pro  ·  R$ 49,90/mês','Proposta por médico: múltiplos convênios, competências disponíveis, metas, regras e importação.','PARA ACOMPANHAR A ROTINA')
text(s,'A separação de recursos, limites, atendimento e cobrança ainda será implementada e aprovada.',.65,6.15,11.8,.5,15,MUTED)
notes(s,'R$ 49,90 é hipótese comercial nossa, não preço aprovado ou ativo na aplicação. Não alterar a tela atual nem afirmar contratação disponível. Integração hospitalar não está incluída automaticamente no plano individual: depende de acordo institucional e validação técnica. Histórico só existirá para períodos com dados carregados; exportação não é prometida.')

# 12
s=base('Convite ao piloto','Vamos validar com quem vai usar.','Proposta: grupo inicial de 5 a 10 médicos, com acompanhamento por 30 dias.')
card(s,.55,2.8,3.95,'Demonstração','Primeiro encontro com dados fictícios e roteiro guiado.','ETAPA 1')
card(s,4.68,2.8,3.95,'Validação controlada','Dados reais somente após aprovação do hospital, autenticação validada e regras conferidas.','ETAPA 2')
card(s,8.81,2.8,3.95,'Decisão conjunta','Medir clareza, divergências, tempo de conferência e interesse no plano Pro.','ETAPA 3')
text(s,'Pergunta central: “Você entende melhor quanto está previsto e o que precisa conferir?”',.65,6.2,11.9,.55,19,GREEN,True)
notes(s,'Trinta dias gratuitos são uma proposta de piloto a aprovar, sem compromisso de implantação em data fixa. Os 30 dias de uso real só começam após os requisitos de segurança e conciliação estarem atendidos. Não anunciar integrações ou descontos garantidos.')

# 13 — appendix
s=base('Apêndice · situação da entrega','Pronto para mostrar. Validações antes de operar.','Use este apêndice na discussão com gestão, TI e responsáveis pelo piloto.')
card(s,.55,2.75,5.98,'Demonstrável hoje','Dashboard em duas competências; filtros; regras; metas; importação; cálculo SQL; interface móvel e desktop.','MVP FUNCIONAL COM DADOS FICTÍCIOS')
card(s,6.75,2.75,5.98,'Antes dos dados reais','Validar Auth/Supabase, homologar regras e competência, autorizar tratamento de dados, implantar e conciliar.','CONDIÇÕES PARA OPERAÇÃO')
text(s,'Registro da entrega: build aprovado · 9 testes SQL/RLS · 4 testes de navegador · sem acesso ao Oracle.',.65,6.2,11.9,.55,15,GREEN,True)
notes(s,'Base: docs/VALIDACAO.md, resultado de 16/10/2026. Não afirmar certificação de segurança, garantia de exatidão hospitalar ou validação em dispositivo físico. Conector MV ausente. Não há fechamento contábil/snapshot histórico; alterações de regras recalculam a previsão.')

# 14
s=base('Apêndice · modelo institucional sugerido','O hospital pode viabilizar o acesso dos médicos.','Uma proposta comercial inicial, condicionada à descoberta técnica e ao escopo contratual.')
box(s,.55,2.8,7.45,3.8,GREEN)
text(s,'LICENÇA INSTITUCIONAL · PROPOSTA',.85,3.1,6.8,.5,13,WHITE,True)
text(s,'R$ 1.490/mês',.85,3.88,6.7,.8,38,WHITE,True,'Bricolage Grotesque')
text(s,'Referência para até 50 médicos em uma instituição.\nImplantação inicial sugerida: R$ 4.900.',.85,5.0,6.7,1.1,20,WHITE)
text(s,'Definir antes de contratar',8.45,3.0,4.2,.7,23,GREEN,True)
text(s,'• Fonte e escopo dos dados\n• Trabalho de integração MV\n• Suporte e disponibilidade\n• Custos de operação e limites',8.45,4.0,4.1,2.3,18,MUTED)
notes(s,'Estes valores são hipóteses para negociação e não tabela de mercado, orçamento fechado ou valuation. Licença institucional e permissões administrativas ainda não foram implementadas. Não somar automaticamente a mensalidade institucional à individual: escolher quem contrata e evitar cobrança dupla. Integrações customizadas, rede hospitalar, migração, SLA e esforço fora do escopo devem ser estimados após a Fase 0.')

# 15
s=base('Apêndice · hipótese comercial','O que validar antes de fixar os preços.','Preço precisa cobrir operação e suporte e fazer sentido para o médico e para o hospital.')
card(s,.55,2.8,3.95,'Disposição a pagar','Apresentar R$ 49,90 e perguntar se o valor é compatível com a utilidade percebida.','VALIDAÇÃO COM MÉDICOS')
card(s,4.68,2.8,3.95,'Custo de servir','Estimar hospedagem, suporte, aquisição, tributos e manutenção de integrações.','VIABILIDADE')
card(s,8.81,2.8,3.95,'Receita ilustrativa','100 assinantes Pro × R$ 49,90 = R$ 4.990/mês de receita bruta.','CENÁRIO, NÃO PROJEÇÃO')
text(s,'Não é lucro ou previsão de vendas. Não há estimativa defensável de valor de venda do software nesta fase.',.65,6.2,11.8,.55,14,MUTED)
notes(s,'Não equiparar o produto a um ERP, prontuário ou sistema de clínica completo. A consulta pública ao iClinic confirma oferta adjacente de gestão/repasse, mas não valida nosso preço. Para avaliar venda integral do software, levantar propriedade intelectual, receita recorrente, contratos, custos, passivos e esforço restante; não inventar um valuation.')

# 16
s=base('Apêndice · referências e premissas','Material de apoio à apresentação.','Base da demonstração, limites conhecidos e contexto comercial consultado.')
text(s,'Produto e demonstração',.65,2.8,6,.5,23,GREEN,True)
text(s,'Plano do projeto e regras do MVP\nRegistro de validação: 16/10/2026\nCapturas reais: outubro/novembro de 2026 (fictícios)',.65,3.45,11.8,1.4,18,MUTED)
text(s,'Referência externa de oferta adjacente',.65,5.05,11.8,.5,23,GREEN,True)
text(s,'iClinic · Gestão financeira e repasse para profissionais\nhttps://site.iclinic.com.br/funcionalidades/gestao-financeira/',.65,5.65,11.8,.8,16,MUTED)
notes(s,'Consulta web realizada nesta preparação, em 16/10/2026. Referência de categoria, não comparativo equivalente nem recomendação de concorrente. Valores de Repasse MED são sugestões próprias ainda não aprovadas. Documentos locais: docs/PLANO.md, docs/REGRAS.md, docs/VALIDACAO.md, docs/FASE-3.md e docs/apresentacao/GUIA-APRESENTADOR.md.')

target=OUT/'Repasse_MED_Apresentacao_Medicos.pptx'
prs.core_properties.title='Repasse MED — apresentação aos médicos'
prs.core_properties.subject='Demonstração funcional e proposta comercial para validação'
prs.core_properties.author='Repasse MED'
prs.core_properties.keywords='MVP, dados fictícios, repasse médico, proposta de planos'
prs.save(target)
print(f'PPTX gerado: {target} ({len(prs.slides)} slides)')
