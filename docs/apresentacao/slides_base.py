"""Generate an editable physician-facing deck. Run with PYTHONPATH=.tools/presentation."""
from pathlib import Path
from decimal import Decimal
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from PIL import ImageDraw
from io import BytesIO

ROOT = Path(__file__).resolve().parents[2]
BRAND = re.search(r'PRODUCT_NAME\s*=\s*"([^"]+)"', (ROOT/'web/src/lib/brand.ts').read_text(encoding='utf-8')).group(1)
_brand_source = (ROOT/'web/src/lib/brand.ts').read_text(encoding='utf-8')
_colors = re.search(r'BRAND_COLORS\s*=\s*\{([^}]+)\}', _brand_source).group(1)
def brand_color(name): return re.search(rf'\b{name}:\s*"#([0-9a-fA-F]{{6}})"', _colors).group(1).upper()
_css = (ROOT/'web/src/app/globals.css').read_text(encoding='utf-8')
_light = re.search(r':root\s*\{([^}]+)\}', _css).group(1)
def ui_color(name):
    value = re.search(rf'--{name}:\s*#([0-9a-fA-F]{{3,6}})', _light).group(1).upper()
    return ''.join(c*2 for c in value) if len(value) == 3 else value
_plans=(ROOT/'web/src/lib/plans.ts').read_text(encoding='utf-8')
PRO_PRICE=Decimal(re.search(r'\bindividual:\s*([\d.]+)',_plans).group(1))
def brl(value): return 'R$ '+f'{value:,.2f}'.replace(',','_').replace('.',',').replace('_','.')
OUT = ROOT / 'docs' / 'apresentacao'
OUT.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG, WHITE, INK, MUTED = brand_color('paper'), ui_color('surface'), ui_color('ink'), ui_color('muted')
GREEN, SOFT, LINE, GOLD, DEEP = ui_color('accent'), ui_color('accent-soft'), ui_color('line'), brand_color('gold'), brand_color('deep')
LIGHT_ON_DARK = brand_color('paper')
# Segoe UI ships with Windows/Office, so the deck looks the same in PowerPoint and in the exported PDF.
BODY, DISPLAY = 'Segoe UI', 'Segoe UI Semibold'
LEFT = .55  # one left edge for every element on the slide
TOTAL = 19
LOGO_PRIMARY = ROOT/'web/public/brand/logo-primary.png'
LOGO_INVERSE = ROOT/'web/public/brand/logo-inverse.png'

def box(s,x,y,w,h,color=WHITE,radius=True,line=None,round_=.12):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: sh.adjustments[0]=round_
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color)
    if line: sh.line.color.rgb=RGBColor.from_string(line); sh.line.width=Pt(1)
    else: sh.line.fill.background()
    sh.shadow.inherit=False  # flat surfaces; the default theme shadow looked muddy on the cards
    return sh

def text(s,content,x,y,w,h,size=20,color=INK,bold=False,font=BODY,spacing=1.1,after=6,align=None):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01); tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.line_spacing=spacing;p.space_after=Pt(after)
        if align:p.alignment=align
    return sh

def base(kicker,title,sub=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(DEEP if dark else BG)
    if title: logo(s,dark,11.84,.28,1.0)
    text(s,kicker.upper(),LEFT,.4,11,.3,10.5,LIGHT_ON_DARK if dark else GREEN,True,after=0)
    text(s,title,LEFT,.82,12.2,1.1,32,WHITE if dark else INK,False,DISPLAY,spacing=1.0,after=0)
    if sub:text(s,sub,LEFT,1.75,12.1,.5,16,LIGHT_ON_DARK if dark else MUTED,after=0)
    text(s,f'{BRAND}  ·  Controle pessoal  ·  capturas com dados fictícios',LEFT,7.08,10,.2,9,LIGHT_ON_DARK if dark else MUTED,after=0)
    text(s,f'{len(prs.slides):02} / {TOTAL}',11.4,7.06,1.4,.25,10,LIGHT_ON_DARK if dark else MUTED,True,after=0,align=3)  # 3 = right
    return s

def card(s,x,y,w,title,body,tag=None,dark=False,h=3.1):
    box(s,x,y,w,h,GREEN if dark else WHITE,line=None if dark else LINE)
    if tag:text(s,tag,x+.28,y+.26,w-.56,.3,11,LIGHT_ON_DARK if dark else GREEN,True,after=0)
    text(s,title,x+.28,y+.7,w-.56,.8,22,WHITE if dark else INK,False,DISPLAY,spacing=1.0,after=0)
    text(s,body,x+.28,y+1.75,w-.56,h-1.9,16,WHITE if dark else MUTED,after=0)

def pill(s,x,y,w,h,label,fill=GREEN,color=WHITE,size=14):
    box(s,x,y,w,h,fill,round_=.5)
    text(s,label,x,y+(h-size/72*1.3)/2,w,h,size,color,True,after=0,align=2)  # 2 = center

def notes(s,body):s.notes_slide.notes_text_frame.text=body

def logo(s,dark,x,y,w):
    path=LOGO_INVERSE if dark else LOGO_PRIMARY
    iw,ih=Image.open(path).size
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(w*ih/iw))

def screenshot(s,path,x,y,w,h):
    image=Image.open(path).convert('RGB')
    iw,ih=image.size
    source=str(path)
    # The archived local login captures contain the demo account and masked password.
    # Remove both field contents from the copy embedded in the public deck.
    if path.name == 'login-390.png':
        draw=ImageDraw.Draw(image)
        draw.rectangle((54,577,334,609), fill='#FFFFFF')
        draw.rectangle((54,659,334,690), fill='#FFFFFF')
    elif path.name == 'login-1280.png':
        draw=ImageDraw.Draw(image)
        draw.rectangle((796,395,1185,425), fill='#FFFFFF')
        draw.rectangle((796,478,1185,509), fill='#FFFFFF')
    if path.name.startswith('login-'):
        source=BytesIO()
        image.save(source,format='PNG')
        source.seek(0)
    # Crop the viewport proportionally; never stretch a browser capture.
    p=s.shapes.add_picture(source, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    source_ratio=iw/ih; frame_ratio=w/h
    if source_ratio<frame_ratio:p.crop_bottom=1-source_ratio/frame_ratio  # keep the top of tall page captures
    else:
        crop=(1-frame_ratio/source_ratio)/2;p.crop_left=crop;p.crop_right=crop
    return p
